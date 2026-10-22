"""Polished, traceable DOCX export for one resident-deliberation session."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = RGBColor(31, 58, 82)
TEAL = RGBColor(12, 124, 134)
INK = RGBColor(34, 34, 34)
MUTED = RGBColor(91, 104, 112)
WHITE = RGBColor(255, 255, 255)
CORE_FIELDS = ["risk_perception", "benefit_perception", "distributive_fairness", "procedural_fairness", "trust", "place_attachment", "social_influence", "conditional_acceptance", "conflict_action_likelihood"]
COMPENSATION_FIELDS = ["electricity_discount", "one_time_cash", "long_term_living_support", "community_fund", "local_participation", "loss_compensation", "formal_council", "monitoring_committee"]
FIELD_LABELS = {"risk_perception": "위험 인식", "benefit_perception": "편익 인식", "distributive_fairness": "분배 공정성", "procedural_fairness": "절차 공정성", "trust": "신뢰", "place_attachment": "장소 애착", "social_influence": "사회적 영향", "conditional_acceptance": "조건부 수용", "conflict_action_likelihood": "갈등행동 가능성", "electricity_discount": "전기요금 감면", "one_time_cash": "일시금", "long_term_living_support": "장기 생활지원", "community_fund": "지역공동기금", "local_participation": "주민참여", "loss_compensation": "손실 보전", "formal_council": "공식 협의체", "monitoring_committee": "주민 감시기구"}


def _duration(seconds: int | None) -> str:
    if seconds is None:
        return "미기록"
    hours, rest = divmod(max(0, int(seconds)), 3600)
    minutes, seconds = divmod(rest, 60)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}"


def _value(value: Any, fallback: str = "미기록") -> str:
    return fallback if value in (None, "") else str(value)


def _set_run_font(run: Any, name: str, size: float, color: RGBColor = INK, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def _shade(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _cell_margin(cell: Any, top: int = 90, start: int = 120, bottom: int = 90, end: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_width(cell: Any, width_dxa: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    width = properties.first_child_found_in("w:tcW")
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:w"), str(width_dxa))
    width.set(qn("w:type"), "dxa")


def _fixed_layout(table: Any, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    properties = table._tbl.tblPr
    layout = properties.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    table_width = properties.first_child_found_in("w:tblW")
    if table_width is not None:
        table_width.set(qn("w:w"), str(sum(widths)))
        table_width.set(qn("w:type"), "dxa")
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            _set_cell_width(cell, widths[index])
            _cell_margin(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _add_page_field(paragraph: Any) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin); run._r.append(instruction); run._r.append(end)
    _set_run_font(run, "Noto Sans KR", 8, MUTED)


def _configure(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)
    normal = document.styles["Normal"]
    normal.font.name = "Noto Serif KR"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Serif KR")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.28
    for style_name, size, color, before, after in (("Title", 25, NAVY, 0, 8), ("Heading 1", 16, NAVY, 18, 8), ("Heading 2", 12.5, TEAL, 12, 5), ("Heading 3", 11, NAVY, 8, 4)):
        style = document.styles[style_name]
        style.font.name = "Noto Sans KR"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans KR")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    bullet = document.styles["List Bullet"]
    bullet.font.name = "Noto Serif KR"
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Serif KR")
    bullet.font.size = Pt(10.5)
    bullet.paragraph_format.space_after = Pt(4)
    bullet.paragraph_format.line_spacing = 1.22
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("UDL PERSONA  |  RESEARCH SESSION REPORT")
    _set_run_font(header_run, "Noto Sans KR", 8, MUTED, True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    prefix = footer.add_run("UDL Persona  |  ")
    _set_run_font(prefix, "Noto Sans KR", 8, MUTED)
    _add_page_field(footer)


def _heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def _paragraph(document: Document, text: str, emphasized: bool = False) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    _set_run_font(run, "Noto Serif KR", 10.5, INK, emphasized)
    paragraph.paragraph_format.line_spacing = 1.28
    paragraph.paragraph_format.space_after = Pt(7)


def _bullets(document: Document, values: list[str], empty: str = "기록이 없습니다.") -> None:
    for value in values or [empty]:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(str(value))
        _set_run_font(run, "Noto Serif KR", 10.5)


def _table(document: Document, headers: list[str], rows: list[list[Any]], widths: list[int] | None = None, compact: bool = False) -> None:
    widths = widths or [9360 // len(headers)] * len(headers)
    widths[-1] += 9360 - sum(widths)
    table = document.add_table(rows=1, cols=len(headers))
    _fixed_layout(table, widths)
    for index, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.text = ""
        _shade(cell, "203A52")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(str(header))
        _set_run_font(run, "Noto Sans KR", 8.5 if compact else 9, WHITE, True)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, row)):
            cell.text = ""
            if row_index % 2 == 1:
                _shade(cell, "F5F8FA")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(_value(value, "-"))
            _set_run_font(run, "Noto Sans KR" if compact else "Noto Serif KR", 8.2 if compact else 9.2)
            if len(str(value)) < 12:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _callout(document: Document, label: str, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    _fixed_layout(table, [9360])
    cell = table.cell(0, 0)
    cell.text = ""
    _shade(cell, "E8F2F3")
    paragraph = cell.paragraphs[0]
    label_run = paragraph.add_run(label + "  ")
    _set_run_font(label_run, "Noto Sans KR", 9, TEAL, True)
    text_run = paragraph.add_run(text)
    _set_run_font(text_run, "Noto Serif KR", 10, INK)
    document.add_paragraph().paragraph_format.space_after = Pt(3)


def _question_rows(analysis: dict[str, Any]) -> list[list[str]]:
    info = analysis.get("input", {})
    recipes, mapping = info.get("automatic_recipe", {}), info.get("mapping", {})
    rows: list[list[str]] = []
    for field in CORE_FIELDS + COMPENSATION_FIELDS:
        parts = recipes.get(field, []) if isinstance(recipes, dict) else []
        if parts:
            sources = "\n".join(str(item.get("header", "")) for item in parts)
            scoring = "역채점 후 평균" if any(item.get("reverse") for item in parts) else "평균"
        else:
            source = mapping.get(field, "") if isinstance(mapping, dict) else ""
            sources = str(source) if source else "미연결"
            scoring = "원문 수치" if source else "중립값 3.0"
        rows.append([FIELD_LABELS[field], "군집 투입" if field in CORE_FIELDS else "보상 선호", scoring, sources])
    return rows


def _survey_method_text(analysis: dict[str, Any]) -> list[str]:
    info = analysis.get("input", {})
    return [
        f"입력 방식은 {_value(info.get('analysis_mode'), '수동 항목 연결')}이다. 위험 인식, 편익 인식, 분배 공정성, 절차 공정성, 신뢰, 장소 애착, 사회적 영향, 조건부 수용, 갈등행동 가능성의 9개 지표를 군집분석 변수로 사용했다.",
        f"문항을 숫자로 변환한 뒤 9개 핵심 지표 중 하나라도 결측인 응답은 완전사례 분석에서 제외했다. 유효 응답은 {_value(info.get('valid_respondents'))}명, 제외 응답은 {_value(info.get('dropped_rows'), '0')}명이다. 결측을 임의 대치하지 않은 대신, 결측의 체계성은 원자료에서 별도로 점검해야 한다.",
        f"입력 척도는 {_value(info.get('input_scale'))}이다. 1~7점 원문은 1 + (원점수-1) x 4/6으로 1~5점 범위에 선형 변환했고, 이미 1~5점인 지표는 유지했다. 자동 연결에서 역문항은 역채점 후 동일 개념 문항의 평균으로 산출했다.",
        "보상 관련 8개 항목은 군집을 나누는 변수가 아니라 군집별 보상·절차 선호를 요약하는 보조지표다. 연결되지 않은 보상 항목의 중립값 3.0은 결측 대체가 아니라 비교 가능한 중립 기준이며, 해당 항목의 군집 간 차이는 해석하지 않는다.",
    ]


def _cluster_method_text(analysis: dict[str, Any]) -> list[str]:
    validation, selected = analysis.get("validation", {}), analysis.get("validation", {}).get("selected_metrics", {})
    return [
        "9개 핵심 지표를 StandardScaler로 표준화한 뒤 유클리드 거리 기반 K-means를 적용했다. 표준화는 특정 문항의 분산이 군집 형성에 과도하게 작용하는 것을 막는다.",
        "설정된 k 후보마다 초기 중심을 50회 바꾸어 적합하고, 실루엣 계수(높을수록 양호), Calinski-Harabasz 지수(높을수록 양호), Davies-Bouldin 지수(낮을수록 양호)의 순위를 합산했다. 순위합이 가장 낮은 k를 선택하고, 동점이면 실루엣이 높은 k와 더 작은 k를 우선했다.",
        f"선택 결과는 k={_value(validation.get('selected_k'))}, 실루엣={_value(selected.get('silhouette'))}, Calinski-Harabasz={_value(selected.get('calinski_harabasz'))}, Davies-Bouldin={_value(selected.get('davies_bouldin'))}이다. 최종 K-means는 100회 초기화로 다시 적합했다.",
        "안정성 점검으로 Ward 연결법 계층적 군집과의 Adjusted Rand Index(ARI), 그리고 복원추출 20회의 재표본 K-means 평균 ARI를 기록했다. 이 지표는 재현성의 보조증거이며, 자연군집의 존재를 확정하는 검정은 아니다.",
    ]


def _panel_rows(discussion: dict[str, Any]) -> list[list[str]]:
    rows = []
    for resident in discussion.get("panel", []):
        background = resident.get("background", {})
        demographic = " / ".join(value for value in (_value(background.get("age"), "미연결"), _value(background.get("sex"), "미연결")) if value)
        rows.append([resident.get("display_name", resident.get("speaker_id", "주민")), f"군집 {resident.get('cluster_id', '-')}", demographic, _value(background.get("occupation"), "미연결"), _value(background.get("district"), "미연결")])
    return rows


def _background_method_text(discussion: dict[str, Any]) -> list[str]:
    panel = discussion.get("panel", [])
    linked = sum(str(item.get("background", {}).get("source", "")).startswith("nvidia/") for item in panel)
    return [
        "군집 소속, 보상 선호, 토론 핵심 조건은 입력 설문으로 결정한다. NVIDIA 배경은 발화의 생활 맥락을 부여하는 보조 자료이며, 연령·성별·직업은 군집 소속이나 수용 판정에 가중하지 않는다.",
        "NVIDIA Nemotron-Personas-Korea에서 선택 지역과 정확히 같은 시·군·구(district) 행을 먼저 조회한다. 정확 일치가 없을 때만 같은 광역권(province) 행으로 보완하고, 연결이 실패하면 가상의 인구특성을 만들지 않고 미연결로 기록한다.",
        f"조회 필드는 age, sex, occupation, district, province, housing_type, family_type, persona, cultural_background, family_persona이며, 이번 세션의 연결 인원은 {linked}/{len(panel)}명이다. 이 공개 페르소나는 실제 해당 지역 주민 표본이나 지역 여론을 뜻하지 않는다.",
    ]


def _representative_remarks(discussion: dict[str, Any], limit: int = 6) -> list[dict[str, Any]]:
    selected, seen = [], set()
    for turn in discussion.get("turns", []):
        speaker = str(turn.get("speaker_id", ""))
        if speaker in seen or not str(turn.get("spoken_text", "")).strip():
            continue
        if turn.get("claim_type") not in {"근거주장", "확인질문", "생활경험", "가치판단"}:
            continue
        selected.append(turn); seen.add(speaker)
        if len(selected) >= limit:
            break
    return selected


def _append_remarks(document: Document, discussion: dict[str, Any]) -> None:
    for item in _representative_remarks(discussion):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(4)
        label = paragraph.add_run(f"{item.get('utterance_id', '-')}  {item.get('speaker_name', '주민')}  |  {item.get('claim_type', '-')} · {item.get('tone', '-')}")
        _set_run_font(label, "Noto Sans KR", 9, TEAL, True)
        text = document.add_paragraph()
        text.paragraph_format.left_indent = Inches(0.16)
        text.paragraph_format.space_after = Pt(8)
        run = text.add_run(str(item.get("spoken_text", "")))
        _set_run_font(run, "Noto Serif KR", 10)


def create_report_docx(output_path: Path, discussion: dict[str, Any], config: dict[str, Any], report: dict[str, Any]) -> Path:
    """Create a professor-facing research report with readable main text and appendices."""
    document = Document(); _configure(document)
    outcome, process = report.get("outcome", {}), report.get("process", {})
    final_confirmation = report.get("final_confirmation", {})
    reproducibility, analysis = report.get("reproducibility", {}), report.get("analysis_snapshot", {})
    facility = config.get("facility_name", "비선호시설")
    region = config.get("region_display") or f"{config.get('province', '')} {config.get('district', '')}".strip() or "지역 미기록"
    kicker = document.add_paragraph(); kicker.paragraph_format.space_before = Pt(28); kicker.paragraph_format.space_after = Pt(10)
    _set_run_font(kicker.add_run("RESEARCH SESSION REPORT"), "Noto Sans KR", 10, TEAL, True)
    title = document.add_paragraph(); title.paragraph_format.space_after = Pt(8)
    _set_run_font(title.add_run(f"{facility} 주민 숙의 시뮬레이션\n연구 보고서"), "Noto Sans KR", 25, NAVY, True)
    subtitle = document.add_paragraph(); subtitle.paragraph_format.space_after = Pt(20)
    _set_run_font(subtitle.add_run(f"대상 지역  {region}    |    생성일  {datetime.now().strftime('%Y-%m-%d %H:%M')}"), "Noto Sans KR", 10, MUTED)
    _callout(document, "보고서의 성격", "설문 기반 페르소나와 AI 숙의의 탐색적 적용 결과입니다. 실제 주민투표, 지역 여론 또는 금전 보상액의 확정 근거로 사용하지 않습니다.")
    _table(document, ["결과", "최고 수용", "발언 수", "중간점검", "경과 시간"], [[outcome.get("label", "분석 대기"), f"{outcome.get('acceptance_count', 0)}/{outcome.get('panel_size', 0)}", process.get("total_turns", 0), process.get("interim_checks", 0), _duration(process.get("elapsed_seconds"))]], [2200, 1600, 1300, 1700, 2560])

    _heading(document, "1. 핵심 결과")
    _paragraph(document, _value(outcome.get("rule"), "결과 분류 정보가 없습니다."), emphasized=True)
    _heading(document, "도출된 수용 조건", level=2)
    _bullets(document, [str(item) for item in outcome.get("package", [])], "도출된 수용 조건이 없습니다.")
    remaining = [f"{item.get('speaker_name', '주민')}: {item.get('requested_change', '수정 요청')}" for item in outcome.get("remaining_requests", []) if item.get("requested_change")]
    if remaining:
        _callout(document, "남은 수정 요청", " / ".join(remaining))
    _heading(document, "잠정 합의문 최종 확인", level=2)
    _paragraph(document, final_confirmation.get("label", "최종 확인 기록이 없습니다."), emphasized=True)
    if final_confirmation.get("responses"):
        _table(document, ["주민", "최종 확인", "발언", "남은 요청"], [[item.get("speaker_name"), item.get("decision"), item.get("spoken_text"), item.get("requested_change")] for item in final_confirmation.get("responses", [])], [1700, 1300, 4800, 1800], compact=True)
    elif outcome.get("classification") == "qualified_consensus":
        _callout(document, "해석 유의", "조건부 합의는 전원 최종 확인 가결이 아닙니다. 사전 등록한 75% 이상 수용·거부 0명 기준을 만족한 중간점검 결과이며, 수정 의견은 잔존할 수 있습니다.")

    document.add_page_break()
    _heading(document, "2. 설문 해석과 군집분석")
    _heading(document, "2.1 설문 데이터 해석 기준", level=2)
    for text in _survey_method_text(analysis): _paragraph(document, text)
    _heading(document, "2.2 군집 수 선정과 검증", level=2)
    for text in _cluster_method_text(analysis): _paragraph(document, text)
    validation = analysis.get("validation", {})
    candidates = [[item.get("k"), item.get("silhouette"), item.get("calinski_harabasz"), item.get("davies_bouldin"), item.get("composite_rank"), "선택" if item.get("k") == validation.get("selected_k") else "후보"] for item in validation.get("all_candidates", [])]
    _table(document, ["k", "실루엣", "CH", "DB", "순위합", "판정"], candidates or [["-", "-", "-", "-", "-", "기록 없음"]], [700, 1500, 2000, 1600, 1500, 2060])
    _callout(document, "재현성 기록", f"난수 시드 {reproducibility.get('seed')} | Ward-K-means ARI {reproducibility.get('ward_ari')} | 20회 재표본 평균 ARI {reproducibility.get('bootstrap_ari')}")

    _heading(document, "3. 페르소나와 주민 패널")
    _paragraph(document, "각 군집의 대표 응답자는 표준화된 9차원 공간에서 군집 중심에 가장 가까운 medoid이다. 12명 패널은 군집 비율을 최대잔여 방식으로 배분하고, 군집 수가 12명 이하이면 모든 군집을 최소 한 명 포함한다. 군집 안에서는 중심 사례, 위험·갈등 성향이 높은 사례, 조건부 수용·절차공정성이 높은 사례, 이후 서로 가장 다른 사례 순으로 선정한다.")
    allocation = [[item.get("cluster_id"), item.get("name"), item.get("survey_members"), item.get("survey_share"), item.get("panel_members")] for item in report.get("panel_allocation", [])]
    _table(document, ["군집", "페르소나", "설문 인원", "비율", "패널"], allocation or [["-", "-", "-", "-", "-"]], [700, 4200, 1500, 1200, 1760])
    _heading(document, "NVIDIA 주민 배경의 적용", level=2)
    for text in _background_method_text(discussion): _paragraph(document, text)
    _table(document, ["주민", "군집", "연령 / 성별", "직업", "배경 지역"], _panel_rows(discussion), [1200, 1000, 1600, 2700, 2860])

    document.add_page_break()
    _heading(document, "4. 숙의 절차와 종료 판정")
    _paragraph(document, "개회 후 모든 패널이 최소 한 번 발언하고, 이후에는 3개 발언마다 사회자가 직전 주장·근거 요청·미해결 쟁점을 정리한다. 수정안은 3~6개 조항으로 제시하며, 초기안에 없는 금액·거리·기간·사고확률은 임의로 만들지 못하게 제한한다. 수정안이 있을 때만 전원에게 수용(accept)·수정(revise)·거부(reject)를 묻는 중간점검을 실시한다. NVIDIA 주민 배경은 연령·성별·직업·주거·가구의 토론용 맥락에만 사용하며, 원본 서사·지역 사실·개인사는 주민 발언의 증거로 사용하지 않도록 제한한다.")
    _paragraph(document, f"사전 설정은 최대 {reproducibility.get('max_turns')}발언, 정체 판정 전 최소 {reproducibility.get('min_turns_before_unresolved')}발언, 조건부 합의 종료 전 최소 {reproducibility.get('min_turns_before_qualified_consensus')}발언이다. 전원 수용이면 잠정 합의문을 다시 읽고 전원 confirm을 받아 만장일치 가결로 종료한다.")
    _callout(document, "75% 수용을 바로 종료하지 않는 이유", "75% 이상 수용과 거부 0명은 조건부 합의의 필요조건이다. 다만 최소 발언 수 전의 수용은 쟁점 미노출이나 초기 다수의 일시적 수용일 수 있어, 사전 설정한 숙의량 뒤에 같은 문안이 유지되는지를 확인한다.")
    _paragraph(document, "최소 발언 수 이후 75% 이상 수용·거부 0명이면 조건부 합의로 종료한다. 거부가 남거나 최근 중간점검에서 수용 인원이 개선되지 않으면 정체에 따른 미가결 종료로 기록한다. 최대 발언 수에 도달하면 가장 높은 수용 수정안을 분석 대상으로만 제시하며, 이를 가결과 혼동하지 않는다.")
    trajectory = [[f"{item.get('check')}차" + (" (분석 채택)" if item.get("is_best") else ""), item.get("accept", 0), item.get("revise", 0), item.get("reject", 0)] for item in report.get("trajectory", [])]
    _table(document, ["중간점검", "수용", "수정", "거부"], trajectory or [["-", "-", "-", "-"]], [3900, 1800, 1800, 1860])

    _heading(document, "5. 회의록 점검과 주요 논의")
    audit = process.get("transcript_audit", {}); summary = audit.get("summary", {})
    _table(document, ["반복 후보", "대화 연결 오류", "비연결 후보", "맥락 불일치 후보", "사회자 경고"], [[summary.get("near_repeat_count", 0), summary.get("invalid_reply_count", 0), summary.get("unmatched_turn_count", 0), summary.get("scenario_mismatch_count", 0), summary.get("generation_warning_count", 0)]], [1900, 1900, 1800, 1900, 1860])
    _paragraph(document, "자동 점검은 문장 유사도, reply_to 형식, 시설·지역 명칭, 구조화 응답 경고를 기반으로 한다. 문맥 적합성이나 설득의 질은 자동으로 확정할 수 없으므로, 원문 회의록의 질적 검토가 필요하다.")
    flags = audit.get("flags", [])
    if flags:
        _heading(document, "우선 검토가 필요한 항목", level=2)
        _table(document, ["유형", "위치", "내용"], [[item.get("type"), item.get("location"), item.get("detail")] for item in flags[:8]], [2000, 1600, 5760], compact=True)
        _paragraph(document, f"자동 점검 전체 {len(flags)}건은 부록 B에 발언 번호별로 수록했다.")
    _heading(document, "주요 쟁점", level=2)
    _table(document, ["쟁점", "언급 횟수"], [[item.get("label"), item.get("count")] for item in report.get("issues", [])], [7000, 2360])
    _heading(document, "대표 발언", level=2); _append_remarks(document, discussion)

    _heading(document, "6. 해석 범위와 재현 정보")
    _paragraph(document, "이 결과는 특정 모델과 시드에서 생성된 한 번의 시뮬레이션 궤적이다. 패널의 수용은 실제 주민 동의율이나 정책 수용성의 모집단 추정치가 아니다. 설문 기반 군집의 상대적 차이와 사전 설정한 합의 규칙에서 어떤 조항이 쟁점이 되었는지를 탐색하는 데 해석 범위를 한정한다. 학술적 활용에서는 원자료 설문표, 문항-지표 연결의 전문가 검토, 다른 k·시드·모델의 민감도 분석, 실제 숙의 또는 후속 검증 자료를 함께 제시해야 한다. 특히 결과에 포함된 생활 언급은 실제 특정 주민의 진술이나 지역 사실을 재현한 것이 아니라, 제공 시나리오와 토론용 배경 맥락 아래 생성된 합성 발언이다.")
    _table(document, ["항목", "기록"], [["시설", reproducibility.get("facility")], ["지역", reproducibility.get("region")], ["모델", reproducibility.get("model")], ["시드 / 반복", f"{reproducibility.get('seed')} / {reproducibility.get('run_group')} #{reproducibility.get('run_index')}"], ["프롬프트·사실경계", f"{reproducibility.get('prompt_version')} / {reproducibility.get('fact_boundary_version')}"], ["설문 원천", reproducibility.get("survey_source")], ["선택 군집 수", reproducibility.get("selected_k")], ["숙의 참고 틀", reproducibility.get("deliberation_method")], ["종료 사유", reproducibility.get("termination_reason")], ["최종 경과 시간", _duration(reproducibility.get("elapsed_seconds"))]], [2400, 6960])
    repeat = report.get("repeat_summary", {})
    _heading(document, "동일 설계 반복 실행 비교", level=2)
    _paragraph(document, "아래 표는 같은 패널·시나리오 설계를 새 시드로 반복한 기록이다. 반복 횟수가 적으면 분포를 일반화할 수 없으므로, 결과의 안정성 확인을 위한 기술적 기록으로만 해석한다.")
    _table(document, ["회차", "시드", "결과", "수용", "발언 수"], [[item.get("run_index"), item.get("seed"), item.get("label"), item.get("acceptance"), item.get("turns")] for item in repeat.get("runs", [])] or [[1, reproducibility.get("seed"), outcome.get("label"), f"{outcome.get('acceptance_count', 0)}/{outcome.get('panel_size', 0)}", process.get("total_turns", 0)]], [1000, 1900, 2900, 1800, 1760])

    document.add_page_break()
    _heading(document, "부록 A. 문항-지표 연결표")
    _paragraph(document, "자동 연결 또는 수동 연결에 사용된 원문 문항을 아래에 보존한다. 이 표는 분석 재현을 위한 자료이며, 본문 해석에서는 핵심지표와 산출 규칙을 우선 참고한다.")
    _table(document, ["지표", "사용", "산출", "연결된 원문 문항"], _question_rows(analysis), [1300, 1200, 1400, 5460], compact=True)
    _heading(document, "부록 B. 회의록 자동 점검 상세")
    _paragraph(document, "아래 목록은 오류를 확정하는 판정이 아니라 사전 정의한 자동 점검 규칙에 따라 검토가 필요한 발언을 모두 표시한 것이다.")
    _table(document, ["유형", "위치", "내용"], [[item.get("type"), item.get("location"), item.get("detail")] for item in flags] or [["특이사항 없음", "-", "사전 정의한 자동 점검 규칙에서 경고가 발견되지 않았습니다."]], [2000, 1600, 5760], compact=True)
    _heading(document, "부록 C. 패널 배경 적용 기록")
    for resident in discussion.get("panel", []):
        background = resident.get("background", {})
        _callout(document, f"{resident.get('display_name', resident.get('speaker_id', '주민'))} | 군집 {resident.get('cluster_id', '-')}", f"연령 {_value(background.get('age'), '미연결')} · 성별 {_value(background.get('sex'), '미연결')} · 직업 {_value(background.get('occupation'), '미연결')} · 배경 지역 {_value(background.get('district'), '미연결')} · 주거 {_value(background.get('housing_type'), '미기록')} · 가구 {_value(background.get('family_type'), '미기록')}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path
