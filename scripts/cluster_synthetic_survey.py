"""Validated clustering and persona construction for the synthetic substation survey."""

from __future__ import annotations

import argparse
import csv
import json
from itertools import combinations
from pathlib import Path
from statistics import mean
from typing import Any

import numpy as np
from openpyxl import load_workbook
from sklearn.cluster import AgglomerativeClustering, KMeans
from sklearn.metrics import adjusted_rand_score, calinski_harabasz_score, davies_bouldin_score, silhouette_score
from sklearn.preprocessing import StandardScaler


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SURVEY = PROJECT_ROOT.parent / "26.06_1차설문_가상1000명.xlsx"
FEATURE_NAMES = [
    "risk_perception",
    "benefit_perception",
    "distributive_fairness",
    "procedural_fairness",
    "trust",
    "place_attachment",
    "social_influence",
    "conditional_acceptance",
    "conflict_action_likelihood",
]
QUESTION_NUMBERS = [11, 12, 13, 14, 15, 16, 17, 28, 29, 30, 31, 32, 33, 34, 35, 36, 37, 39, 40, 41, 42, 43, 44, 45, 46, 47, 48, 49, 50, 51, 52, 61, 62, 64, 65]
COMPENSATION_ITEMS = {
    "electricity_discount": 44,
    "one_time_cash": 45,
    "long_term_living_support": 46,
    "community_fund": 47,
    "local_participation": 48,
    "loss_compensation": 49,
    "formal_council": 50,
    "monitoring_committee": 51,
}
SCALE_ITEM_SETS = {
    "risk_perception": [31, 32, 65],
    "benefit_perception": [17, 29, 30],
    "procedural_fairness": [28, 37, 64],
    "trust": [34, 35, 36, 62],
    "place_attachment": [11, 12, 13, 14, 15, 16],
    "social_influence": [41, 42, 61],
    "conditional_acceptance": [40, 43, 52],
    "conflict_action_likelihood": [31, 39, 65],
}


def progress(percent: int, message: str) -> None:
    print(f"[{percent:>3}%] {message}", flush=True)


def find_question_column(headers: list[str], number: int) -> int:
    prefix = f"{number}."
    for index, header in enumerate(headers):
        if header.startswith(prefix):
            return index
    raise ValueError(f"{number}번 문항을 찾지 못했습니다.")


def as_likert(value: object, number: int) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"{number}번 문항의 값이 올바르지 않습니다: {value!r}") from error
    if not 1 <= result <= 7:
        raise ValueError(f"{number}번 문항은 1~7 범위여야 합니다: {result}")
    return result


def to_state(values: list[float], reverse: bool = False) -> float:
    transformed = [8 - value if reverse else value for value in values]
    return round(1 + (mean(transformed) - 1) * 4 / 6, 3)


def cronbach_alpha(item_matrix: np.ndarray) -> float | None:
    """Return alpha for multi-item operational composites; None for zero variance."""
    item_count = item_matrix.shape[1]
    if item_count < 2:
        return None
    item_variance = item_matrix.var(axis=0, ddof=1).sum()
    total_variance = item_matrix.sum(axis=1).var(ddof=1)
    if total_variance == 0:
        return None
    return float(item_count / (item_count - 1) * (1 - item_variance / total_variance))


def build_features(row: list[object], columns: dict[int, int]) -> dict[str, float]:
    value = lambda number: as_likert(row[columns[number]], number)
    return {
        "risk_perception": to_state([value(31), value(32), value(65)]),
        "benefit_perception": to_state([value(17), value(29), value(30)]),
        "distributive_fairness": to_state([value(33)], reverse=True),
        "procedural_fairness": to_state([value(28), value(37), value(64)]),
        "trust": to_state([value(34), value(35), value(36), value(62)]),
        "place_attachment": to_state([value(11), value(12), value(13), value(14), value(15), value(16)]),
        "social_influence": to_state([value(41), value(42), value(61)]),
        "conditional_acceptance": to_state([value(40), value(43), value(52)]),
        "conflict_action_likelihood": to_state([value(31), value(39), value(65)]),
    }


def load_survey(path: Path) -> tuple[list[dict[str, Any]], dict[int, int], dict[str, float | None]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["가상 응답 1000명"]
    rows = list(worksheet.iter_rows(values_only=True))
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    columns = {number: find_question_column(headers, number) for number in QUESTION_NUMBERS}
    id_column = headers.index("가상 응답자 ID")
    agents = []
    for row in rows[1:]:
        features = build_features(list(row), columns)
        compensation = {
            name: round(1 + (as_likert(row[columns[number]], number) - 1) * 4 / 6, 3)
            for name, number in COMPENSATION_ITEMS.items()
        }
        agents.append({"agent_id": str(row[id_column]), "features": features, "compensation_preferences": compensation})
    reliability = {}
    for scale_name, item_numbers in SCALE_ITEM_SETS.items():
        item_matrix = np.array(
            [[as_likert(row[columns[number]], number) for number in item_numbers] for row in rows[1:]], dtype=float
        )
        alpha = cronbach_alpha(item_matrix)
        reliability[scale_name] = None if alpha is None else round(alpha, 4)
    return agents, columns, reliability


def rank_desc(values: list[float]) -> list[int]:
    order = sorted(range(len(values)), key=lambda index: values[index], reverse=True)
    ranks = [0] * len(values)
    for rank, index in enumerate(order, start=1):
        ranks[index] = rank
    return ranks


def rank_asc(values: list[float]) -> list[int]:
    order = sorted(range(len(values)), key=lambda index: values[index])
    ranks = [0] * len(values)
    for rank, index in enumerate(order, start=1):
        ranks[index] = rank
    return ranks


def assess_candidates(matrix: np.ndarray, seed: int, minimum_k: int, maximum_k: int) -> list[dict[str, Any]]:
    candidates = []
    for k in range(minimum_k, maximum_k + 1):
        model = KMeans(n_clusters=k, n_init=100, random_state=seed)
        labels = model.fit_predict(matrix)
        candidates.append(
            {
                "k": k,
                "silhouette": float(silhouette_score(matrix, labels)),
                "calinski_harabasz": float(calinski_harabasz_score(matrix, labels)),
                "davies_bouldin": float(davies_bouldin_score(matrix, labels)),
                "inertia": float(model.inertia_),
            }
        )
    silhouette_ranks = rank_desc([item["silhouette"] for item in candidates])
    ch_ranks = rank_desc([item["calinski_harabasz"] for item in candidates])
    db_ranks = rank_asc([item["davies_bouldin"] for item in candidates])
    for index, item in enumerate(candidates):
        item["composite_rank"] = silhouette_ranks[index] + ch_ranks[index] + db_ranks[index]
    return candidates


def bootstrap_stability(matrix: np.ndarray, k: int, seed: int, repetitions: int = 30) -> float:
    random = np.random.default_rng(seed)
    predictions: list[np.ndarray] = []
    sample_size = round(matrix.shape[0] * 0.8)
    for _ in range(repetitions):
        indexes = random.choice(matrix.shape[0], size=sample_size, replace=False)
        model = KMeans(n_clusters=k, n_init=50, random_state=int(random.integers(1, 2**31 - 1)))
        model.fit(matrix[indexes])
        predictions.append(model.predict(matrix))
    scores = [adjusted_rand_score(left, right) for left, right in combinations(predictions, 2)]
    return float(np.mean(scores))


def persona_name(profile: dict[str, float], global_mean: dict[str, float]) -> str:
    if profile["risk_perception"] >= global_mean["risk_perception"] + 0.25 and profile["trust"] <= global_mean["trust"] - 0.15:
        return "안전검증 우선형"
    if profile["distributive_fairness"] <= global_mean["distributive_fairness"] - 0.2 or profile["loss_compensation"] >= 3.7:
        return "형평·피해보전 민감형"
    if profile["benefit_perception"] >= global_mean["benefit_perception"] + 0.2 and profile["risk_perception"] <= global_mean["risk_perception"]:
        return "공공필요 수용형"
    if profile["procedural_fairness"] >= global_mean["procedural_fairness"] + 0.15:
        return "참여협의 우선형"
    return "신중검토형"


def profile_summary(profile: dict[str, float], global_mean: dict[str, float]) -> str:
    high = [name for name in FEATURE_NAMES if profile[name] >= global_mean[name] + 0.2]
    low = [name for name in FEATURE_NAMES if profile[name] <= global_mean[name] - 0.2]
    high_text = ", ".join(high) if high else "뚜렷한 상대적 고점 없음"
    low_text = ", ".join(low) if low else "뚜렷한 상대적 저점 없음"
    return f"상대적으로 높은 지표: {high_text}. 상대적으로 낮은 지표: {low_text}."


def main() -> None:
    parser = argparse.ArgumentParser(description="가상 설문 기반 변전소 페르소나 군집분석")
    parser.add_argument("--survey", type=Path, default=DEFAULT_SURVEY)
    parser.add_argument("--min-k", type=int, default=2)
    parser.add_argument("--max-k", type=int, default=8)
    parser.add_argument("--seed", type=int, default=2606211000)
    args = parser.parse_args()
    if args.min_k < 2 or args.max_k < args.min_k:
        raise ValueError("군집 수 범위는 2 이상이며 min-k <= max-k 이어야 합니다.")

    progress(5, "1,000명 가상 설문 읽는 중")
    agents, _, reliability = load_survey(args.survey)
    progress(18, "9개 이론 기반 지표 표준화 중")
    raw_matrix = np.array([[agent["features"][name] for name in FEATURE_NAMES] for agent in agents], dtype=float)
    scaled_matrix = StandardScaler().fit_transform(raw_matrix)

    progress(30, f"K-means 후보 비교 중: k={args.min_k}~{args.max_k}")
    candidates = assess_candidates(scaled_matrix, args.seed, args.min_k, args.max_k)
    selected = min(candidates, key=lambda item: (item["composite_rank"], -item["silhouette"], item["k"]))
    selected_k = int(selected["k"])

    progress(52, f"재표본 안정성 검증 중: k={selected_k}")
    stability = bootstrap_stability(scaled_matrix, selected_k, args.seed)
    progress(68, "계층적 Ward 군집과의 일치도 검증 중")
    kmeans = KMeans(n_clusters=selected_k, n_init=100, random_state=args.seed)
    labels = kmeans.fit_predict(scaled_matrix)
    ward_labels = AgglomerativeClustering(n_clusters=selected_k, linkage="ward").fit_predict(scaled_matrix)
    method_agreement = float(adjusted_rand_score(labels, ward_labels))

    # K-means labels are arbitrary. Order them by increasing conditional acceptance
    # so the persona identifiers stay interpretable across reruns using this seed.
    label_order = sorted(
        range(selected_k),
        key=lambda cluster: (
            raw_matrix[labels == cluster, FEATURE_NAMES.index("conditional_acceptance")].mean(),
            -raw_matrix[labels == cluster, FEATURE_NAMES.index("conflict_action_likelihood")].mean(),
        ),
    )
    label_map = {old_label: new_label for new_label, old_label in enumerate(label_order)}
    labels = np.array([label_map[int(label)] for label in labels], dtype=int)

    progress(78, "군집 대표 페르소나 구성 중")
    global_mean = {name: float(raw_matrix[:, index].mean()) for index, name in enumerate(FEATURE_NAMES)}
    personas = []
    clustered_rows = []
    for cluster_id in range(selected_k):
        indexes = np.where(labels == cluster_id)[0]
        centroid = scaled_matrix[indexes].mean(axis=0)
        distances = np.linalg.norm(scaled_matrix[indexes] - centroid, axis=1)
        medoid_index = int(indexes[int(np.argmin(distances))])
        feature_profile = {name: round(float(raw_matrix[indexes, column].mean()), 3) for column, name in enumerate(FEATURE_NAMES)}
        compensation_profile = {
            name: round(float(mean(agents[index]["compensation_preferences"][name] for index in indexes)), 3)
            for name in COMPENSATION_ITEMS
        }
        combined_profile = {**feature_profile, **compensation_profile}
        label = persona_name(combined_profile, {**global_mean, **{name: 3.0 for name in COMPENSATION_ITEMS}})
        personas.append(
            {
                "cluster_id": int(cluster_id + 1),
                "persona_name": f"페르소나 {cluster_id + 1}: {label}",
                "member_count": int(len(indexes)),
                "member_share": round(float(len(indexes) / len(agents)), 4),
                "representative_agent_id": agents[medoid_index]["agent_id"],
                "state_profile": feature_profile,
                "compensation_preferences": compensation_profile,
                "profile_summary": profile_summary(feature_profile, global_mean),
            }
        )
        for index in indexes:
            clustered_rows.append({"agent_id": agents[index]["agent_id"], "cluster_id": int(cluster_id + 1), **agents[index]["features"], **agents[index]["compensation_preferences"]})
    personas.sort(key=lambda persona: persona["cluster_id"])
    clustered_rows.sort(key=lambda row: row["agent_id"])

    progress(92, "검증 결과와 페르소나 파일 저장 중")
    reports_dir = PROJECT_ROOT / "reports"
    agents_dir = PROJECT_ROOT / "agents"
    reports_dir.mkdir(exist_ok=True)
    agents_dir.mkdir(exist_ok=True)
    with (reports_dir / "cluster_selection_metrics.csv").open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(candidates[0].keys()))
        writer.writeheader()
        writer.writerows(candidates)
    with (agents_dir / "clustered_synthetic_agents.csv").open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(clustered_rows[0].keys()))
        writer.writeheader()
        writer.writerows(clustered_rows)
    validation = {
        "input": {"synthetic_respondents": len(agents), "feature_names": FEATURE_NAMES, "seed": args.seed},
        "operational_scale_reliability": {
            "cronbach_alpha": reliability,
            "note": "단일문항 형평성 지표는 알파를 산출하지 않았으며, 일부 문항은 이론적으로 중복될 수 있어 탐색적 구성개념 점검으로 해석합니다.",
        },
        "selection_method": "K-means candidates ranked by silhouette, Calinski-Harabasz, and Davies-Bouldin criteria",
        "selected_k": selected_k,
        "selected_metrics": selected,
        "all_candidates": candidates,
        "bootstrap_mean_ari": round(stability, 4),
        "ward_kmeans_adjusted_rand_index": round(method_agreement, 4),
        "interpretation": "Silhouette values below 0.20 should be interpreted as exploratory segmentation rather than strongly separated natural clusters.",
    }
    (reports_dir / "cluster_validation.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8")
    (agents_dir / "clustered_personas.json").write_text(json.dumps({"personas": personas}, ensure_ascii=False, indent=2), encoding="utf-8")
    progress(100, f"완료: 선택 군집 수 k={selected_k}, 페르소나 {len(personas)}개")


if __name__ == "__main__":
    main()
