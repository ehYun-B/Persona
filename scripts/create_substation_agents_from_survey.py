"""Map the synthetic survey workbook to substation deliberation agent profiles."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path
from statistics import mean

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SURVEY = PROJECT_ROOT.parent / "26.06_1차설문_가상1000명.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "agents" / "generated_substation_agents.json"
DEFAULT_SUMMARY = PROJECT_ROOT / "reports" / "substation_agent_summary.json"


def report_progress(percent: int, message: str) -> None:
    print(f"[{percent:>3}%] {message}", flush=True)


def question_column(headers: list[str], number: int) -> int:
    prefix = f"{number}."
    for index, header in enumerate(headers):
        if str(header).strip().startswith(prefix):
            return index
    raise ValueError(f"설문에서 {number}번 문항을 찾을 수 없습니다.")


def as_likert(value: object, question_number: int) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"{question_number}번 문항의 리커트 값이 비어 있거나 올바르지 않습니다: {value!r}") from error
    if not 1 <= number <= 7:
        raise ValueError(f"{question_number}번 문항 값은 1~7이어야 합니다: {number}")
    return number


def to_state(values: list[float], reverse: bool = False) -> float:
    scores = [8 - value if reverse else value for value in values]
    return round(1 + (mean(scores) - 1) * 4 / 6, 2)


def classify_agent(state: dict[str, float]) -> str:
    if state["risk_perception"] >= 3.7 and state["trust"] < 3.0:
        return "안전우려·불신형"
    if state["conditional_acceptance"] >= 3.5 and state["procedural_fairness"] >= 3.4:
        return "조건협의형"
    if state["benefit_perception"] >= 3.5 and state["risk_perception"] <= 2.8:
        return "공공필요·수용형"
    return "신중검토형"


def build_state(row: list[object], columns: dict[int, int]) -> dict[str, float]:
    value = lambda number: as_likert(row[columns[number]], number)
    return {
        "risk_perception": to_state([value(31), value(32), value(65)]),
        "benefit_perception": to_state([value(17), value(29), value(30)]),
        "distributive_fairness": to_state([value(33)], reverse=True),
        "procedural_fairness": to_state([value(28), value(37), value(64)]),
        "trust": to_state([value(34), value(35), value(36), value(62)]),
        "place_attachment": to_state([value(11), value(12), value(13), value(14), value(15), value(16)]),
        "social_influence": to_state([value(41), value(42), value(61)]),
        "conditional_acceptance": to_state([value(40), value(43), value(52)]),
        "conflict_action_likelihood": to_state([value(31), value(39), value(65)]),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="가상 설문을 변전소 토론용 주민 에이전트로 변환합니다.")
    parser.add_argument("--survey", type=Path, default=DEFAULT_SURVEY)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--summary", type=Path, default=DEFAULT_SUMMARY)
    args = parser.parse_args()

    report_progress(5, "가상 설문 파일 확인 중")
    if not args.survey.exists():
        raise FileNotFoundError(f"가상 설문 파일을 찾지 못했습니다: {args.survey}")

    report_progress(15, "엑셀 시트 읽는 중")
    workbook = load_workbook(args.survey, read_only=True, data_only=True)
    if "가상 응답 1000명" not in workbook.sheetnames:
        raise ValueError("'가상 응답 1000명' 시트가 필요합니다.")
    worksheet = workbook["가상 응답 1000명"]
    rows = list(worksheet.iter_rows(values_only=True))
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    columns = {number: question_column(headers, number) for number in [11, 12, 13, 14, 15, 16, 17, 28, 29, 30, 31, 32, 33, 34, 35, 36, 37, 39, 40, 41, 42, 43, 52, 61, 62, 64, 65]}

    report_progress(30, f"변전소 태도 지표 계산 중: {len(rows) - 1}명")
    agents: list[dict[str, object]] = []
    source_id_index = headers.index("가상 응답자 ID")
    for index, row in enumerate(rows[1:], start=1):
        state = build_state(list(row), columns)
        agents.append(
            {
                "agent_id": str(row[source_id_index]),
                "cluster_name": classify_agent(state),
                "facility_focus": "substation",
                "state": state,
                "provenance": "synthetic_survey_1000",
            }
        )
        if index % 100 == 0 or index == len(rows) - 1:
            report_progress(30 + round(index / (len(rows) - 1) * 55), f"에이전트 변환 완료: {index}/{len(rows) - 1}")

    report_progress(90, "요약 통계와 파일 저장 중")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.summary.parent.mkdir(parents=True, exist_ok=True)
    cluster_counts = Counter(str(agent["cluster_name"]) for agent in agents)
    dimension_means = {
        dimension: round(mean(float(agent["state"][dimension]) for agent in agents), 3)
        for dimension in agents[0]["state"]
    }
    args.output.write_text(
        json.dumps({"metadata": {"facility_focus": "substation", "agent_count": len(agents)}, "agents": agents}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    args.summary.write_text(
        json.dumps(
            {"agent_count": len(agents), "cluster_counts": cluster_counts, "dimension_means": dimension_means},
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    report_progress(100, f"완료: {args.output}")


if __name__ == "__main__":
    main()
